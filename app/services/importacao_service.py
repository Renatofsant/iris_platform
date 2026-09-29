"""
Importação de materiais do professor (PDF, DOCX, TXT) para o campo de texto da adaptação.

O arquivo é lido em memória e descartado; só o texto extraído volta ao navegador, para o
professor revisar antes de adaptar. PDFs digitalizados (só imagem) não têm texto: nesses
casos a orientação é usar o modo "Imagem ou equação".
"""

from __future__ import annotations

import io
import logging
import re
import zipfile

from app.services.ai_service import TAMANHO_MAXIMO_TEXTO

logger = logging.getLogger(__name__)

MAX_PAGINAS_PDF = 40
FORMATOS = {".pdf": "PDF", ".docx": "Word", ".txt": "texto"}


class ImportacaoInvalidaError(ValueError):
    def __init__(self, mensagem: str, status: int = 400):
        super().__init__(mensagem)
        self.mensagem = mensagem
        self.status = status


def _formato(nome: str, conteudo: bytes) -> str:
    """Decide pelo conteúdo (assinatura do arquivo), usando a extensão só para o texto puro."""
    if conteudo.startswith(b"%PDF-"):
        return ".pdf"
    if conteudo.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(conteudo)) as pacote:
                if "word/document.xml" in pacote.namelist():
                    return ".docx"
        except zipfile.BadZipFile:
            pass
        raise ImportacaoInvalidaError("Arquivo compactado não reconhecido. Envie PDF, DOCX ou TXT.")
    if nome.lower().endswith(".txt"):
        return ".txt"
    raise ImportacaoInvalidaError("Formato não suportado. Envie um arquivo PDF, DOCX ou TXT.")


def _limpar(texto: str) -> str:
    texto = texto.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    texto = re.sub(r"(\w)-\n(\w)", r"\1\2", texto)  # palavras hifenizadas na quebra de linha do PDF
    texto = re.sub(r"[ \t]+", " ", texto)
    texto = re.sub(r" *\n *", "\n", texto)
    return re.sub(r"\n{3,}", "\n\n", texto).strip()


def extrair_texto(nome: str, conteudo: bytes) -> dict:
    if not conteudo:
        raise ImportacaoInvalidaError("O arquivo está vazio.")
    formato = _formato(nome or "", conteudo)
    avisos: list[str] = []
    paginas = None

    try:
        if formato == ".pdf":
            from pypdf import PdfReader

            leitor = PdfReader(io.BytesIO(conteudo))
            if leitor.is_encrypted:
                raise ImportacaoInvalidaError("O PDF está protegido por senha.")
            paginas = len(leitor.pages)
            if paginas > MAX_PAGINAS_PDF:
                avisos.append(f"Foram lidas só as primeiras {MAX_PAGINAS_PDF} de {paginas} páginas.")
            texto = "\n\n".join((pagina.extract_text() or "") for pagina in leitor.pages[:MAX_PAGINAS_PDF])
        elif formato == ".docx":
            import docx

            documento = docx.Document(io.BytesIO(conteudo))
            partes = []
            for paragrafo in documento.paragraphs:
                estilo = (paragrafo.style.name or "").lower() if paragrafo.style is not None else ""
                nivel = re.search(r"(?:heading|título)\s*(\d)", estilo)
                prefixo = "#" * min(int(nivel.group(1)), 3) + " " if nivel and paragrafo.text.strip() else ""
                partes.append(prefixo + paragrafo.text)
            for tabela in documento.tables:  # tabelas viram Markdown, que a adaptação entende
                linhas = [[c.text.strip() for c in linha.cells] for linha in tabela.rows]
                if linhas:
                    partes.append("\n| " + " | ".join(linhas[0]) + " |\n|" + "---|" * len(linhas[0]))
                    partes.extend("| " + " | ".join(l) + " |" for l in linhas[1:])
            texto = "\n".join(partes)
        else:
            try:
                texto = conteudo.decode("utf-8-sig")
            except UnicodeDecodeError:
                texto = conteudo.decode("latin-1")
    except ImportacaoInvalidaError:
        raise
    except Exception as exc:
        logger.warning("Falha ao ler arquivo importado (%s): %s", formato, exc)
        raise ImportacaoInvalidaError("Não foi possível ler o arquivo. Ele pode estar corrompido.") from exc

    texto = _limpar(texto)
    if len(texto) < 20:
        if formato == ".pdf":
            raise ImportacaoInvalidaError(
                "Este PDF não tem texto selecionável (parece digitalizado). "
                "Use o modo \"Imagem ou equação\" com uma captura da página.", status=422,
            )
        raise ImportacaoInvalidaError("O arquivo não tem texto suficiente para adaptar.", status=422)
    if len(texto) > TAMANHO_MAXIMO_TEXTO:
        texto = texto[:TAMANHO_MAXIMO_TEXTO].rsplit("\n", 1)[0]
        avisos.append(f"O texto foi cortado em {TAMANHO_MAXIMO_TEXTO:,} caracteres. Divida o material em partes.".replace(",", "."))
    return {"texto": texto, "formato": FORMATOS[formato], "paginas": paginas, "caracteres": len(texto), "avisos": avisos}
