"""
Exportação de provas e atividades adaptadas para Word (.docx) com python-docx.

O material salvo (HTML já sanitizado) é convertido em estilos reais do Word:
  • fonte e tamanho escolhidos pelo professor aplicados ao estilo "Normal" (em pt);
    títulos escalam a partir dele, então o aluno de baixa visão recebe exatamente 28 pt;
  • entrelinhas, margens A4 e idioma pt-BR (para corretor e leitores de tela);
  • cabeçalho escolar em tabela; campos vazios viram linhas para preencher à mão;
  • checklist com ☐, passos numerados, glossário, tabelas e pictogramas ARASAAC.
O documento é editável: o professor pode ajustá-lo no Word antes de imprimir.

Limitação: fórmulas saem como código LaTeX (seguido da leitura por extenso), não como
equações nativas do Word.
"""

from __future__ import annotations

import io
import logging
import re
import unicodedata
import urllib.request
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.services.impressao_service import ORIGEM_IMAGENS_PERMITIDA

logger = logging.getLogger(__name__)

FONTES_WORD = {
    "padrao": "Arial",  # Inter raramente está instalada no Word; Arial é universal e legível
    "atkinson": "Atkinson Hyperlegible",
    "dyslexic": "OpenDyslexic",
}
ESCALA_TITULOS = {1: 1.3, 2: 1.18, 3: 1.08, 4: 1.0, 5: 1.0, 6: 1.0}
PRETO = RGBColor(0, 0, 0)
TAGS_VAZIAS = {"br", "img", "hr"}
BLOCOS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "li", "dl", "dt", "dd", "table",
          "section", "div", "header", "figure", "figcaption", "blockquote", "pre", "hr"}


# ---------------------------------------------------------------------------
# Árvore HTML mínima (o HTML já foi sanitizado pelo nh3 ao salvar)
# ---------------------------------------------------------------------------

@dataclass
class No:
    tag: str
    attrs: dict[str, str] = field(default_factory=dict)
    filhos: list[Any] = field(default_factory=list)

    @property
    def classes(self) -> set[str]:
        return set(self.attrs.get("class", "").split())

    def texto(self) -> str:
        return "".join(f if isinstance(f, str) else f.texto() for f in self.filhos)


class _Construtor(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.raiz = No("raiz")
        self.pilha = [self.raiz]

    def handle_starttag(self, tag, attrs):
        no = No(tag, {k: v or "" for k, v in attrs})
        self.pilha[-1].filhos.append(no)
        if tag not in TAGS_VAZIAS:
            self.pilha.append(no)

    def handle_endtag(self, tag):
        for i in range(len(self.pilha) - 1, 0, -1):
            if self.pilha[i].tag == tag:
                del self.pilha[i:]
                break

    def handle_data(self, data):
        self.pilha[-1].filhos.append(data)


def _arvore(html: str) -> No:
    construtor = _Construtor()
    construtor.feed(html)
    construtor.close()
    return construtor.raiz


# ---------------------------------------------------------------------------
# Utilidades de formatação
# ---------------------------------------------------------------------------

def _definir_fonte(estilo_ou_run, nome: str) -> None:
    estilo_ou_run.font.name = nome
    rpr = estilo_ou_run.element.get_or_add_rPr()
    fontes = rpr.find(qn("w:rFonts"))
    if fontes is None:
        fontes = OxmlElement("w:rFonts")
        rpr.append(fontes)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fontes.set(qn(atributo), nome)


def _idioma_pt_br(estilo) -> None:
    rpr = estilo.element.get_or_add_rPr()
    idioma = OxmlElement("w:lang")
    idioma.set(qn("w:val"), "pt-BR")
    rpr.append(idioma)


def _borda_paragrafo(paragrafo, lados=("top", "left", "bottom", "right"), espessura_oitavos=12) -> None:
    ppr = paragrafo._p.get_or_add_pPr()
    bordas = OxmlElement("w:pBdr")
    for lado in lados:
        borda = OxmlElement(f"w:{lado}")
        borda.set(qn("w:val"), "single")
        borda.set(qn("w:sz"), str(espessura_oitavos))  # em oitavos de ponto
        borda.set(qn("w:space"), "6")
        borda.set(qn("w:color"), "000000")
        bordas.append(borda)
    ppr.append(bordas)


def _manter_com_proximo(paragrafo) -> None:
    paragrafo.paragraph_format.keep_with_next = True


def _meio_ponto(valor: float) -> Pt:
    """O Word guarda tamanhos de fonte em meios-pontos e o python-docx trunca: arredonda antes."""
    return Pt(round(valor * 2) / 2)


def _espacos(texto: str) -> str:
    return re.sub(r"\s+", " ", texto)


def nome_arquivo(titulo: str) -> str:
    base = unicodedata.normalize("NFKD", titulo).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9]+", "-", base).strip("-").lower()[:60]
    return f"{base or 'material-adaptado'}.docx"


# ---------------------------------------------------------------------------
# Conversor
# ---------------------------------------------------------------------------

class ConversorDocx:
    def __init__(self, preferencias: dict[str, Any], baixar_imagens: bool = True):
        self.prefs = preferencias
        self.fonte_pt = float(preferencias["fonte_pt"])
        self.nome_fonte = FONTES_WORD.get(preferencias.get("tipografia"), "Arial")
        self.baixar_imagens = baixar_imagens
        self._cache_imagens: dict[str, bytes | None] = {}
        self.doc = Document()
        self._configurar_documento()

    # -- configuração ------------------------------------------------------

    def _configurar_documento(self) -> None:
        secao = self.doc.sections[0]
        secao.page_width, secao.page_height = Mm(210), Mm(297)
        secao.top_margin, secao.bottom_margin = Mm(15), Mm(16)
        secao.left_margin = secao.right_margin = Mm(14)

        normal = self.doc.styles["Normal"]
        _definir_fonte(normal, self.nome_fonte)
        _idioma_pt_br(normal)
        normal.font.size = _meio_ponto(self.fonte_pt)
        normal.font.color.rgb = PRETO
        normal.paragraph_format.line_spacing = float(self.prefs.get("entrelinhas", 1.5))
        normal.paragraph_format.space_after = Pt(round(self.fonte_pt * 0.45, 1))
        normal.paragraph_format.widow_control = True

        for nivel, escala in ESCALA_TITULOS.items():
            estilo = self.doc.styles[f"Heading {nivel}"]
            _definir_fonte(estilo, self.nome_fonte)
            estilo.font.size = _meio_ponto(self.fonte_pt * escala)
            estilo.font.bold = True
            estilo.font.italic = False
            estilo.font.color.rgb = PRETO  # o padrão do Word é azul: no P&B tudo é preto
            estilo.paragraph_format.space_before = Pt(round(self.fonte_pt * 0.8, 1))
            estilo.paragraph_format.space_after = Pt(round(self.fonte_pt * 0.3, 1))
            estilo.paragraph_format.keep_with_next = True

    # -- API -----------------------------------------------------------------

    def gerar(self, titulo: str, html: str, cabecalho: dict[str, str], guia: dict[str, Any] | None = None) -> bytes:
        propriedades = self.doc.core_properties
        propriedades.title = titulo
        propriedades.language = "pt-BR"
        propriedades.author = "Plataforma Íris"

        self._cabecalho_escolar(cabecalho)
        self.doc.add_heading(titulo, level=1)
        for filho in _arvore(html).filhos:
            self._bloco(filho, nivel_lista=0)
        self._linhas_resposta(int(self.prefs.get("linhas_resposta", 0)))
        if guia:
            self._guia_mediador(guia)

        saida = io.BytesIO()
        self.doc.save(saida)
        return saida.getvalue()

    # -- cabeçalho, respostas e guia -----------------------------------------------

    def _cabecalho_escolar(self, cab: dict[str, str]) -> None:
        tamanho = _meio_ponto(min(max(11, self.fonte_pt * 0.62), 18))
        linhas = [
            [("Escola", cab.get("escola"))],
            [("Aluno(a)", cab.get("aluno"))],
            [("Turma", cab.get("turma")), ("Data", cab.get("data"))],
            [("Disciplina", cab.get("disciplina")), ("Professor(a)", cab.get("professor"))],
        ]
        tabela = self.doc.add_table(rows=len(linhas), cols=2)
        tabela.style = "Table Grid"
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, campos in enumerate(linhas):
            celulas = tabela.rows[i].cells
            if len(campos) == 1:
                celulas = [celulas[0].merge(celulas[1])]
            for celula, (rotulo, valor) in zip(celulas, campos):
                paragrafo = celula.paragraphs[0]
                paragrafo.paragraph_format.space_after = Pt(2)
                paragrafo.paragraph_format.line_spacing = 1.2
                r = paragrafo.add_run(f"{rotulo}: ")
                r.bold, r.font.size = True, tamanho
                r = paragrafo.add_run((valor or "").strip() or "_" * 28)  # vazio: linha para escrever
                r.font.size = tamanho
        self.doc.add_paragraph()

    def _linhas_resposta(self, total: int) -> None:
        if total <= 0:
            return
        self.doc.add_heading("Respostas", level=2)
        for _ in range(total):
            paragrafo = self.doc.add_paragraph()
            paragrafo.paragraph_format.space_after = Pt(0)
            paragrafo.paragraph_format.space_before = Pt(round(self.fonte_pt * 0.9, 1))
            _borda_paragrafo(paragrafo, lados=("bottom",), espessura_oitavos=8)

    def _guia_mediador(self, guia: dict[str, Any]) -> None:
        self.doc.add_page_break()
        self.doc.add_heading("Orientações para o Mediador/Professor", level=2)
        nota = self.doc.add_paragraph("Página para o professor — não entregar ao estudante.")
        nota.runs[0].italic = True
        nomes = {"elementar": "Elementar", "basico": "Básico", "intermediario": "Intermediário", "avancado": "Avançado"}
        if guia.get("nivel_original") or guia.get("nivel_adaptado"):
            p = self.doc.add_paragraph()
            p.add_run("Complexidade: ").bold = True
            p.add_run(f"{nomes.get(guia.get('nivel_original'), '—')} → {nomes.get(guia.get('nivel_adaptado'), '—')}")
            if guia.get("justificativa"):
                p.add_run(f". {guia['justificativa']}")
        for chave, rotulo in (("ajustes_realizados", "Ajustes realizados"),
                              ("sugestoes_mediacao", "Sugestões de mediação"),
                              ("pontos_de_atencao", "Pontos de atenção")):
            itens = guia.get(chave) or []
            if itens:
                self.doc.add_heading(rotulo, level=3)
                for item in itens:
                    self._paragrafo_lista("• ", 1).add_run(item)

    # -- blocos ------------------------------------------------------------

    def _paragrafo_lista(self, marcador: str, nivel: int, negrito_marcador: bool = False):
        paragrafo = self.doc.add_paragraph()
        recuo = Pt(self.fonte_pt * 1.4)
        paragrafo.paragraph_format.left_indent = recuo * nivel
        paragrafo.paragraph_format.first_line_indent = -recuo
        paragrafo.paragraph_format.space_after = Pt(round(self.fonte_pt * 0.25, 1))
        marca = paragrafo.add_run(marcador)
        marca.bold = negrito_marcador
        return paragrafo

    def _bloco(self, no: Any, nivel_lista: int) -> None:
        if isinstance(no, str):
            if no.strip():
                self._inline(self.doc.add_paragraph(), no, {})
            return
        tag, classes = no.tag, no.classes

        if "iris-guia" in classes:
            return  # o guia é gerado a partir dos dados estruturados, em página própria
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            titulo = self.doc.add_heading("", level=min(int(tag[1]), 6))
            self._inline_filhos(titulo, no, {})
        elif tag in ("p", "figcaption", "dt", "dd", "pre", "blockquote"):
            paragrafo = self.doc.add_paragraph()
            if tag == "dt":
                _manter_com_proximo(paragrafo)
                self._inline_filhos(paragrafo, no, {"bold": True})
                return
            if tag in ("dd", "blockquote"):
                paragrafo.paragraph_format.left_indent = Pt(self.fonte_pt * 1.2)
            if "iris-resumo" in classes:
                _borda_paragrafo(paragrafo)
            self._inline_filhos(paragrafo, no, {"code": tag == "pre"})
        elif tag == "ul" and "iris-ilustrado" in classes:
            for item in (f for f in no.filhos if isinstance(f, No) and f.tag == "li"):
                self._cartao_ilustrado(item)
        elif tag in ("ul", "ol"):
            self._lista(no, nivel_lista + 1)
        elif tag == "table":
            self._tabela(no)
        elif tag == "img":
            self._imagem(no)
        elif tag == "hr":
            _borda_paragrafo(self.doc.add_paragraph(), lados=("bottom",), espessura_oitavos=6)
        elif "iris-pagina-professor" in classes:
            self.doc.add_page_break()  # gabarito: página própria, para o professor
            for filho in no.filhos:
                self._bloco(filho, nivel_lista)
        else:  # section, div, header, figure, dl, span solto...
            filhos_inline = [f for f in no.filhos if isinstance(f, str) or f.tag not in BLOCOS | {"img"}]
            if filhos_inline and all(isinstance(f, str) or f.tag not in BLOCOS for f in no.filhos) \
                    and "".join(f if isinstance(f, str) else f.texto() for f in filhos_inline).strip():
                paragrafo = self.doc.add_paragraph()
                self._inline_filhos(paragrafo, no, {})
                return
            for filho in no.filhos:
                self._bloco(filho, nivel_lista)

    def _lista(self, lista: No, nivel: int) -> None:
        ordenada = lista.tag == "ol"
        checklist = "iris-checklist" in lista.classes
        alternativas = "iris-alternativas" in lista.classes
        numero = int(lista.attrs.get("start") or 1)
        for item in (f for f in lista.filhos if isinstance(f, No) and f.tag == "li"):
            marcador = "☐ " if checklist else "○ " if alternativas else (f"{numero}. " if ordenada else "• ")
            numero += 1
            paragrafo = self._paragrafo_lista(marcador, nivel, negrito_marcador=ordenada)
            self._conteudo_item(item, paragrafo, nivel)

    def _conteudo_item(self, item: No, paragrafo, nivel: int) -> None:
        """Conteúdo de <li>: inline no parágrafo do marcador; blocos seguintes, recuados."""
        atual, usado = paragrafo, False
        for filho in item.filhos:
            if isinstance(filho, str) or filho.tag not in BLOCOS:
                self._inline(atual, filho, {})
                usado = usado or bool((filho if isinstance(filho, str) else filho.texto()).strip())
            elif filho.tag in ("ul", "ol"):
                self._lista(filho, nivel + 1)
            elif filho.tag == "div":
                self._conteudo_item(filho, atual, nivel)
                usado = True
            else:  # <p> dentro do item
                if usado:
                    atual = self.doc.add_paragraph()
                    atual.paragraph_format.left_indent = Pt(self.fonte_pt * 1.4) * nivel
                negrito = "iris-passo-titulo" in filho.classes
                self._inline_filhos(atual, filho, {"bold": negrito})
                usado = True

    def _tabela(self, no: No) -> None:
        linhas = [tr for tr in self._descendentes(no, "tr")]
        if not linhas:
            return
        colunas = max(len([c for c in tr.filhos if isinstance(c, No) and c.tag in ("td", "th")]) for tr in linhas)
        tabela = self.doc.add_table(rows=len(linhas), cols=colunas)
        tabela.style = "Table Grid"
        for i, tr in enumerate(linhas):
            celulas = [c for c in tr.filhos if isinstance(c, No) and c.tag in ("td", "th")]
            for j, celula_html in enumerate(celulas[:colunas]):
                paragrafo = tabela.rows[i].cells[j].paragraphs[0]
                self._inline_filhos(paragrafo, celula_html, {"bold": celula_html.tag == "th"})
        self.doc.add_paragraph()

    def _cartao_ilustrado(self, item: No) -> None:
        tabela = self.doc.add_table(rows=1, cols=2)
        tabela.style = "Table Grid"
        largura_picto = Pt(self.fonte_pt * 4.5)
        tabela.columns[0].width = largura_picto + Mm(4)
        esquerda, direita = tabela.rows[0].cells
        esquerda.width = largura_picto + Mm(4)
        imagem = next(self._descendentes(item, "img"), None)
        dados = self._baixar(imagem.attrs.get("src", "")) if imagem is not None else None
        if dados:
            esquerda.paragraphs[0].add_run().add_picture(io.BytesIO(dados), width=largura_picto)
        texto = next((d for d in self._descendentes(item, "div") if "iris-cartao-texto" in d.classes), item)
        primeiro = True
        for p in (f for f in texto.filhos if isinstance(f, No)):
            paragrafo = direita.paragraphs[0] if primeiro else direita.add_paragraph()
            primeiro = False
            self._inline_filhos(paragrafo, p, {"bold": "iris-conceito" in p.classes})
        self.doc.add_paragraph()

    def _imagem(self, no: No) -> None:
        dados = self._baixar(no.attrs.get("src", ""))
        if dados:
            self.doc.add_paragraph().add_run().add_picture(io.BytesIO(dados), width=Pt(self.fonte_pt * 4.5))
        elif no.attrs.get("alt"):
            self.doc.add_paragraph(f"[Imagem: {no.attrs['alt']}]")

    def _baixar(self, url: str) -> bytes | None:
        """Só pictogramas do ARASAAC (mesma lista permitida na sanitização); falha → sem imagem."""
        if not self.baixar_imagens or not url.startswith(ORIGEM_IMAGENS_PERMITIDA):
            return None
        if url not in self._cache_imagens:
            try:
                with urllib.request.urlopen(url, timeout=5) as resposta:  # noqa: S310 (URL restrita acima)
                    self._cache_imagens[url] = resposta.read(2_000_000)
            except Exception as exc:
                logger.warning("Pictograma não baixado (%s): %s", url, exc)
                self._cache_imagens[url] = None
        return self._cache_imagens[url]

    @staticmethod
    def _descendentes(no: No, tag: str):
        for filho in no.filhos:
            if isinstance(filho, No):
                if filho.tag == tag:
                    yield filho
                yield from ConversorDocx._descendentes(filho, tag)

    # -- inline ------------------------------------------------------------

    def _inline_filhos(self, paragrafo, no: No, estilo: dict[str, bool]) -> None:
        for filho in no.filhos:
            self._inline(paragrafo, filho, estilo)

    def _inline(self, paragrafo, no: Any, estilo: dict[str, bool]) -> None:
        if isinstance(no, str):
            texto = no if estilo.get("code") else _espacos(no)
            if not paragrafo.runs and not estilo.get("code"):
                texto = texto.lstrip()
            if texto:
                run = paragrafo.add_run(texto)
                run.bold = estilo.get("bold") or None
                run.italic = estilo.get("italic") or None
                run.underline = estilo.get("underline") or None
                if estilo.get("code"):
                    _definir_fonte(run, "Consolas")
                if estilo.get("sub"):
                    run.font.subscript = True
                if estilo.get("sup"):
                    run.font.superscript = True
            return
        tag, classes = no.tag, no.classes
        if tag == "br":
            paragrafo.add_run().add_break()
            return
        if tag == "img":
            return
        if classes & {"iris-caixa", "iris-bolinha"}:
            return  # o marcador (☐ ou ○) já foi escrito pela lista
        novo = dict(estilo)
        if tag in ("strong", "b"):
            novo["bold"] = True
        elif tag in ("em", "i"):
            novo["italic"] = True
        elif tag == "u":
            novo["underline"] = True
        elif tag == "code":
            novo["code"] = True
        elif tag == "sub":
            novo["sub"] = True
        elif tag == "sup":
            novo["sup"] = True
        if "iris-math" in classes:
            # \\[ F = m a \\] → F = m a (código LaTeX; a leitura por extenso vem logo abaixo)
            latex = re.sub(r"^\\[\[(]|\\[\])]$", "", no.texto().strip())
            run = paragrafo.add_run(latex)
            _definir_fonte(run, "Cambria Math")
            return
        self._inline_filhos(paragrafo, no, novo)
        if tag == "a" and no.attrs.get("href", "").startswith("http"):
            paragrafo.add_run(f" ({no.attrs['href']})")


def gerar_docx(
    titulo: str,
    html: str,
    preferencias: dict[str, Any],
    cabecalho: dict[str, str],
    guia: dict[str, Any] | None = None,
    baixar_imagens: bool = True,
) -> bytes:
    return ConversorDocx(preferencias, baixar_imagens=baixar_imagens).gerar(titulo, html, cabecalho, guia)
